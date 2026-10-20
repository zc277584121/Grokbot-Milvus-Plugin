"""Real Milvus backends, with one owned collection prefix per test run."""

import os
import uuid
from types import SimpleNamespace

import pytest
from mvbot_common import connect


@pytest.fixture(scope="session")
def client(tmp_path_factory):
    uri = os.getenv("TEST_MILVUS_URI") or str(tmp_path_factory.mktemp("milvus") / "test.db")
    c = connect(SimpleNamespace(uri=uri, token=os.getenv("TEST_MILVUS_TOKEN", ""), db="default"))
    yield c
    c.close()


@pytest.fixture
def collection(client):
    name = "grokbot_test_" + uuid.uuid4().hex[:16]
    assert not client.has_collection(name)
    yield name
    if client.has_collection(name):
        client.drop_collection(name)


class ProbeEmbedder:
    """Deterministic vectors only for retry/identity/score regression tests."""

    provider = "openai"
    model = "probe"

    def __init__(self, *args, **kwargs):
        pass

    def dim(self):
        return 2

    def embed(self, texts):
        return [[1.0, 0.0] if "ZXQERROR42" in t else [0.0, 1.0] for t in texts]


@pytest.fixture
def probe_embedding(monkeypatch):
    import milvus_kb
    import mvbot_common

    monkeypatch.setattr(milvus_kb, "Embedder", ProbeEmbedder)
    monkeypatch.setattr(mvbot_common, "Embedder", ProbeEmbedder)
    for name in ("EMBED_PROVIDER", "EMBED_MODEL", "EMBED_ENDPOINT"):
        monkeypatch.delenv(name, raising=False)


@pytest.fixture
def document_files(tmp_path):
    """Small multilingual corpus with known headings and a page-two citation."""
    import docx
    from fpdf import FPDF

    (tmp_path / "guide.md").write_text(
        "# Error reference\n\n## Disk usage\n\nZXQERROR42 indicates a full disk. Remove unused files to resolve it."
    )
    (tmp_path / "notes.txt").write_text(
        "Search notes\n\nRRF combines ranked lists. The default RRFRanker k is 60."
    )
    doc = docx.Document()
    doc.core_properties.title = "知识库使用说明"
    doc.add_heading("文档更新", level=1)
    doc.add_paragraph("修改文件后重新导入。同一文件使用稳定的文档标识，旧的多余分块会被删除。")
    doc.save(tmp_path / "guide_zh.docx")
    pdf = FPDF()
    pdf.set_title("Document import guide")
    for title, body in [
        (
            "Importing",
            "Extractable text is required for PDF import. Scan-only pages need optical character recognition.",
        ),
        ("Citations", "CITATIONCHECK42 links a result to the original source, heading, and page number."),
    ]:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, title + "\n\n" + body)
    pdf.output(tmp_path / "guide_en.pdf")
    return tmp_path
