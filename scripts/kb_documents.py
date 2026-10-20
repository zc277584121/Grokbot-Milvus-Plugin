"""Document parsing and section-aware chunking, independent of Milvus and embedding services."""

import os
import re

CJK = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")


def _paras(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


def parse_markdown(text, fallback_title):
    blocks, stack, buf, in_code, title = [], [], [], False, None

    def flush():
        if buf:
            body = "\n".join(buf).strip()
            if body:
                blocks.append((" > ".join(h for _, h in stack), None, body))
            buf.clear()

    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            buf.append(line)
            continue
        m = None if in_code else re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if m:
            flush()
            level, h = len(m.group(1)), m.group(2).strip()
            if level == 1 and title is None:
                title = h
            stack[:] = [(level_value, x) for level_value, x in stack if level_value < level] + [(level, h)]
        elif not line.strip() and not in_code:
            flush()
        else:
            buf.append(line)
    flush()
    return title or fallback_title, blocks


def parse_txt(text, fallback_title):
    ps = _paras(text)
    if len(ps) > 1 and "\n" not in ps[0] and len(ps[0]) <= 120:  # short first line = title
        return ps[0], [("", None, p) for p in ps[1:]]
    return fallback_title, [("", None, p) for p in ps]


def parse_pdf(path, fallback_title):
    from pypdf import PdfReader

    r = PdfReader(path)
    title = (r.metadata.title if r.metadata and r.metadata.title else None) or fallback_title
    blocks = []
    for i, page in enumerate(r.pages, 1):
        t = page.extract_text() or ""
        ps = _paras(t) or [level_value.strip() for level_value in t.splitlines() if level_value.strip()]
        blocks += [("", i, re.sub(r"[ \t]*\n[ \t]*", "\n", p)) for p in ps]
    return title, blocks


def parse_docx(path, fallback_title):
    import docx

    d = docx.Document(path)
    title = (d.core_properties.title or "").strip() or None
    blocks, stack = [], []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        m = re.match(r"^(Heading|标题)\s*(\d)", style)
        if style == "Title":
            title = title or t
            continue
        if m:
            level = int(m.group(2))
            if level == 1 and title is None:
                title = t
            stack[:] = [(level_value, x) for level_value, x in stack if level_value < level] + [(level, t)]
            continue
        blocks.append((" > ".join(h for _, h in stack), None, t))
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if rows:
            blocks.append((f"table {ti}", None, "\n".join(rows)))
    return title or fallback_title, blocks


def parse_file(path):
    ext = os.path.splitext(path)[1].lower()
    fallback = os.path.splitext(os.path.basename(path))[0]
    if ext == ".pdf":
        return parse_pdf(path, fallback)
    if ext == ".docx":
        return parse_docx(path, fallback)
    text = open(path, encoding="utf-8", errors="replace").read()
    if ext in (".md", ".markdown"):
        return parse_markdown(text, fallback)
    return parse_txt(text, fallback)


_SENT = re.compile(r"(?<=[。！？!?；;])|(?<=[.!?])\s+")


def _split_long(text, size):
    sents = [s for s in _SENT.split(text) if s and s.strip()]
    out, cur = [], ""
    for s in sents:
        while len(s) > size:  # hard split very long sentences
            if cur:
                out.append(cur)
                cur = ""
            out.append(s[:size])
            s = s[size:]
        if len(cur) + len(s) > size and cur:
            out.append(cur)
            cur = ""
        sep = (
            "" if (not cur or cur.endswith((" ", "\n")) or CJK.search(cur[-1:]) or CJK.search(s[:1])) else " "
        )
        cur += sep + s
    if cur.strip():
        out.append(cur)
    return out


def _tail(text, overlap):
    if overlap <= 0 or len(text) <= overlap:
        return text if overlap > 0 else ""
    t = text[-overlap:]
    m = re.search(r"[。！？!?；;.]\s*", t)  # start the overlap at a sentence boundary when possible
    return t[m.end() :] if m and m.end() < len(t) - 10 else t


def chunk_blocks(blocks, size, overlap):
    """blocks: [(heading, page, text)] -> [(heading, page, chunk_text)]; never crosses a section."""
    chunks, i = [], 0
    while i < len(blocks):
        heading, page = blocks[i][0], blocks[i][1]
        sec = []
        while i < len(blocks) and blocks[i][0] == heading and blocks[i][1] == page:
            sec.append(blocks[i][2])
            i += 1
        pieces = []
        for p in sec:
            lim = max(50, size - overlap)  # leave room for the overlap prefix
            pieces += _split_long(p, lim) if len(p) > lim else [p]
        cur = ""
        for p in pieces:
            if cur and len(cur) + len(p) + 1 > size:
                chunks.append((heading, page, cur))
                tail = _tail(cur, overlap)
                cur = (tail + "\n" if tail else "") + p
            else:
                cur = (cur + "\n" + p) if cur else p
        if cur.strip():
            chunks.append((heading, page, cur))
    return chunks


def lang_of(text):
    if not text:
        return "und"
    return "zh" if len(CJK.findall(text)) / max(1, len(text)) > 0.15 else "en"


def clip_bytes(s, n):
    b = s.encode("utf-8")
    return s if len(b) <= n else b[:n].decode("utf-8", errors="ignore")
